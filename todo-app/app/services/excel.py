"""Bounded Excel parsing shared by read-only uploads and transactional imports."""
from contextlib import contextmanager
from datetime import date, datetime, time, timedelta
from io import BytesIO
from math import isfinite
from typing import BinaryIO
from zipfile import ZipFile

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from pydantic import ValidationError

from app.schemas.excel import ExcelRead, ExcelRow
from app.schemas.todo import TodoWrite

MAX_UPLOAD_BYTES = 2 * 1024 * 1024
MAX_ROWS = 1000
MAX_COLUMNS = 50
MAX_CELL_LENGTH = 5000


def read_upload(filename: str | None, stream: BinaryIO) -> bytes:
    if not (filename or "").lower().endswith(".xlsx"):
        raise HTTPException(415, "Upload an .xlsx file")
    content = stream.read(MAX_UPLOAD_BYTES + 1)
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(413, "Upload must not exceed 2 MiB")
    return content


@contextmanager
def open_workbook(content: bytes):
    workbook = None
    try:
        if len(content) > MAX_UPLOAD_BYTES:
            raise HTTPException(413, "Upload must not exceed 2 MiB")
        with ZipFile(BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > 100 or sum(x.file_size for x in members) > 20 * 1024 * 1024:
                raise HTTPException(413, "Workbook expands beyond allowed limits")
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        if len(workbook.worksheets) != 1:
            raise HTTPException(422, "Use exactly one worksheet")
        sheet = workbook.worksheets[0]
        sheet.reset_dimensions()
        yield sheet
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, "Invalid or unsupported Excel workbook") from exc
    finally:
        if workbook is not None:
            workbook.close()


def read_workbook(content: bytes) -> ExcelRead:
    with open_workbook(content) as sheet:
        iterator = sheet.iter_rows()
        header = list(next(iterator, ()))
        # Excel may retain formatting on otherwise empty trailing columns.
        while header and header[-1].value is None:
            header.pop()
        if not header:
            raise HTTPException(422, "The first row must contain column headers")
        if len(header) > MAX_COLUMNS:
            raise HTTPException(413, f"At most {MAX_COLUMNS} columns are allowed")
        if any(c.data_type in {"f", "e"} or not isinstance(c.value, str)
               or not c.value.strip() or len(c.value) > 200 for c in header):
            raise HTTPException(422, "Headers must be nonempty text, at most 200 characters")
        columns = [c.value.strip() for c in header]
        if len(set(columns)) != len(columns):
            raise HTTPException(422, "Column headers must be unique")
        rows = []
        for number, cells in enumerate(iterator, start=2):
            if number > MAX_ROWS + 1:
                raise HTTPException(413, f"At most {MAX_ROWS} data rows are allowed")
            if all(c.value is None for c in cells):
                continue
            if any(c.value is not None for c in cells[len(columns):]):
                raise HTTPException(422, f"Row {number}: data has no column header")
            values = {}
            for index, name in enumerate(columns):
                cell = cells[index] if index < len(cells) else None
                value = cell.value if cell else None
                if cell and cell.data_type in {"f", "e"}:
                    raise HTTPException(422, f"Row {number}: formulas and Excel errors are not allowed")
                if isinstance(value, str) and len(value) > MAX_CELL_LENGTH:
                    raise HTTPException(422, f"Row {number}: cell exceeds {MAX_CELL_LENGTH} characters")
                if isinstance(value, (datetime, date, time)):
                    value = value.isoformat()
                elif isinstance(value, timedelta):
                    value = str(value)
                if isinstance(value, float) and not isfinite(value):
                    raise HTTPException(422, f"Row {number}: non-finite number")
                values[name] = value
            rows.append(ExcelRow(row_number=number, values=values))
        return ExcelRead(sheet=sheet.title, columns=columns, total=len(rows), rows=rows)


def parse_workbook(content: bytes) -> list[TodoWrite]:
    """Apply todo-specific validation before the caller starts database writes."""
    data = read_workbook(content)
    if data.columns != ["title", "description", "completed"]:
        raise HTTPException(422, "Headers must be: title, description, completed")
    todos, errors = [], []
    for row in data.rows:
        values = dict(row.values)
        if values["completed"] is None:
            values["completed"] = False
        try:
            todos.append(TodoWrite(**values))
        except ValidationError as exc:
            errors.append({"row": row.row_number, "message": "; ".join(e["msg"] for e in exc.errors())})
    if errors:
        raise HTTPException(422, {"message": "Nothing imported; fix these rows", "errors": errors})
    if not todos:
        raise HTTPException(422, "Workbook contains no todos")
    return todos


def create_template() -> bytes:
    workbook = Workbook()
    try:
        workbook.active.append(["title", "description", "completed"])
        workbook.active.append(["Learn FastAPI", "Try authenticated CRUD", False])
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()
    finally:
        workbook.close()

